"""Turns selected cards (or the AI draft) into Markdown or Word files, in English or Bangla."""
import datetime
import re
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

import citations
import sources
from translations import FLAGS, PRIORITIES, TIERS, TYPES, gap_type, method_label
from translations import a as msg

BANGLA = re.compile(r"[ঀ-৿]")
# Nirmala UI ships with Windows 8+ and renders Bangla; Word falls back to Vrinda or other Bangla fonts if missing.
BANGLA_FONT = "Nirmala UI"


def _gap_lines(gaps, lang, by_id):
    L = lambda k, **kw: msg(lang, k, **kw)  # noqa: E731
    out = []
    for n, g in enumerate(gaps, 1):
        out += [f"### {n}. {g['title']}", "",
                f"**{L('g_priority')}:** {PRIORITIES[lang].get(g['priority'], g['priority'])} · **{L('g_type')}:** {gap_type(lang, g['type'])}", ""]
        if g.get("description"):
            out += [g["description"], ""]
        if g.get("signal"):
            out += [f"**{L('g_signal')}:** {g['signal']}", ""]
        ev = [by_id[i]["in_text"] for i in g.get("evidence", []) if i in by_id]
        if ev:
            out += [f"**{L('g_evidence')}:** {', '.join(ev)}", ""]
        if g.get("research_question"):
            out += [f"**{L('g_rq')}:** {g['research_question']}", ""]
        if g.get("method"):
            out += [f"**{L('g_method')}:** {g['method']}", ""]
    return out


def gaps_markdown(run, gaps_result, items):
    """A stand-alone 'research gaps brief' for proposals and 'future research' sections."""
    lang = run.get("lang") or "en"
    L = lambda k, **kw: msg(lang, k, **kw)  # noqa: E731
    by_id = {it["id"]: it for it in items}
    st = gaps_result.get("stats", {})
    gaps = gaps_result.get("gaps", [])
    today = datetime.date.today().strftime("%d %B %Y")
    out = [f"# {L('gb_title')}: {run['topic']}", "", L("gb_generated", d=today, n=st.get("n_relevant", len(items))), "",
           f"> {L('gb_note')}", "", f"## {L('gb_snapshot')}", ""]
    cov = st.get("coverage", [])
    if cov:
        out += [f"- **{c['question']}**: {L('gb_cov', n=c['n'], h=c['high'], f=c['full'])}"
                + (f" · **{L('cov_thin')}**" if c.get("thin") else "") for c in cov] + [""]
    if st.get("methods"):
        out += [f"**{L('gb_methods')}:** " + ", ".join(f"{method_label(lang, m)} ({n})" for m, n in st["methods"]), ""]
    if st.get("countries"):
        out += [f"**{L('gb_countries')}:** " + ", ".join(f"{c} ({n})" for c, n in st["countries"][:10]), ""]
    y = st.get("years") or {}
    if y.get("n"):
        out += [f"**{L('gb_years')}:** {y['oldest']}-{y['newest']}, median {y['median']}, {y['recent']}/{y['n']} ≥ {y['recent_from']}", ""]
    out += [f"## {L('gb_gaps')}", ""] + _gap_lines(gaps, lang, by_id)
    rqs = [g["research_question"] for g in gaps if g.get("research_question")]
    if rqs:
        out += [f"## {L('gb_rqs')}", ""] + [f"{i}. {q}" for i, q in enumerate(rqs, 1)] + [""]
    cited = {i for g in gaps for i in g.get("evidence", []) if i in by_id}
    if cited:
        out += [f"## {L('references')}", ""] + sorted({by_id[i]["citation"] for i in cited}, key=str.lower)
    return "\n".join(out) + "\n"


def evidence_markdown(run, items, plan_questions, gaps_result=None):
    lang = run.get("lang") or "en"
    L = lambda k, **kw: msg(lang, k, **kw)  # noqa: E731
    tier_name = lambda t: TIERS.get(lang, {}).get(t, t)  # noqa: E731
    type_name = lambda t: TYPES.get(lang, {}).get(t) or citations.TYPE_LABELS.get(t, t)  # noqa: E731
    today = datetime.date.today().strftime("%d %B %Y")
    out = [f"# {L('pack_title')}: {run['topic']}", "", L("generated", d=today, n=len(items)), ""]
    if plan_questions:
        out += [f"## {L('questions')}", ""] + [f"- {q}" for q in plan_questions] + [""]
    if gaps_result and gaps_result.get("gaps"):
        out += [f"## {L('gb_gaps')}", "", f"> {L('gb_note')}", ""]
        out += _gap_lines(gaps_result["gaps"], lang, {it["id"]: it for it in items if it.get("id") is not None})
    groups = {}
    for it in items:
        groups.setdefault(it["question"] or "General", []).append(it)
    for question, group in groups.items():
        out += [f"## {question}", ""]
        for it in group:
            src = it.get("via") or ", ".join(sources.label(s) for s in it.get("found_in") or [it["source"]])
            meta = f"**{L('type')}:** {type_name(it['type'])} · **{L('reliability')}:** {tier_name(it['tier'])} · **{L('found_in')}:** {src}"
            if it.get("citations"):
                meta += f" · **{L('cited_by')}:** {it['citations']}"
            if it.get("relevance") is not None:
                meta += f" · **{L('relevance')}:** {it['relevance']}/5"
            out += [f"### {it['title']}", "", meta, ""]
            if it.get("flags"):
                out += ["**⚠ " + " · ".join(FLAGS.get(lang, FLAGS["en"]).get(f, f) for f in it["flags"]) + "**", ""]
            if it.get("future"):
                out += [f"**{L('m_future')}:** {it['future']}", ""]
            for key in ("method", "sample", "limitations"):
                if it.get(key):
                    out += [f"**{L('m_' + {'method': 'method', 'sample': 'sample', 'limitations': 'limits'}[key])}:** {it[key]}", ""]
            if it["findings"]:
                out += [f"**{L('findings')}:**", ""] + [f"- {f}" for f in it["findings"]] + [""]
            if it["data_points"]:
                out += [f"**{L('data')}:**", ""] + [f"- {d}" for d in it["data_points"]] + [""]
            if it.get("quote"):
                out += [f"> \"{it['quote']}\"", ""]
            out += [f"**{L('cite_as')}:** {it['in_text']}", "", f"**{L('reference')}:** {it['citation']}", ""]
            if it.get("url"):
                out += [f"**{L('link')}:** {it['url']}", ""]
    out += [f"## {L('references')}", ""] + sorted({it["citation"] for it in items}, key=str.lower)
    return "\n".join(out) + "\n"


def _set_fonts(style, latin="Calibri"):
    style.font.name = latin
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi"):
        fonts.set(qn(attr), latin)
    fonts.set(qn("w:cs"), BANGLA_FONT)  # Bangla is a "complex script" in Word
    fonts.set(qn("w:eastAsia"), latin)
    rpr_lang = rpr.find(qn("w:lang"))
    if rpr_lang is None:
        rpr_lang = rpr.makeelement(qn("w:lang"), {})
        rpr.append(rpr_lang)
    rpr_lang.set(qn("w:bidi"), "bn-BD")
    if style.font.size:  # Bangla uses the complex-script size, so match it to the normal size
        sz_cs = rpr.find(qn("w:szCs"))
        if sz_cs is None:  # Word requires szCs right after sz and before lang
            sz_cs = rpr.makeelement(qn("w:szCs"), {})
            sz = rpr.find(qn("w:sz"))
            sz.addnext(sz_cs) if sz is not None else rpr_lang.addprevious(sz_cs)
        sz_cs.set(qn("w:val"), str(int(style.font.size.pt * 2)))


def _add_runs(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = p.add_run(part[1:-1])
            run.italic = True
        else:
            run = p.add_run(part)
        if BANGLA.search(part):
            run.font.cs_bold = run.bold  # keep bold/italic on Bangla text
            run.font.cs_italic = run.italic
            run.font.complex_script = True


def md_to_docx(md):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "Quote"):
        try:
            _set_fonts(doc.styles[name])
        except KeyError:
            pass
    in_refs = False
    for line in md.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            heading = m.group(2).replace("*", "")
            doc.add_heading(heading, level=0 if len(m.group(1)) == 1 else len(m.group(1)) - 1)
            in_refs = heading.lower().startswith(("references", "reference list", "তথ্যসূত্র"))
            continue
        if re.match(r"^\s*[-*]\s+", s):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", s))
        elif re.match(r"^\s*\d+\.\s+", s):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", s))
        elif s.startswith(">"):
            _add_runs(doc.add_paragraph(style="Quote"), s.lstrip("> "))
        else:
            p = doc.add_paragraph()
            if in_refs:  # Harvard reference list uses a hanging indent
                p.paragraph_format.left_indent = Pt(36)
                p.paragraph_format.first_line_indent = Pt(-36)
            _add_runs(p, s)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
